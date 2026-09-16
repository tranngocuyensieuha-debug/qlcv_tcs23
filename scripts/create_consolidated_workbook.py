import sys
from pathlib import Path
from shutil import copyfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, column_index_from_string

SCRIPTS_DIR = Path(__file__).resolve().parent
sys.path.append(str(SCRIPTS_DIR.parent))
import seed_data
import server
from scripts.create_web_import_template import create_template, HEADERS, style_range

OUTPUT_FILE = SCRIPTS_DIR.parent.parent / "Bảng tổng hợp đẩy dữ liệu công việc.xlsx"
TEMPLATE_FILE = SCRIPTS_DIR.parent.parent / "mau-du-lieu-day-len-web.xlsx"


def hash_num(s: str, min_val: int, max_val: int) -> int:
    h = 0
    for char in s:
        h = (h * 31 + ord(char)) & 0xFFFFFFFF
    return min_val + (h % (max_val - min_val + 1))


def metric_for(key: str, task_id: str, is_revenue: bool, idx: int = 0):
    base = hash_num(key + task_id, 18, 145) + idx * 3
    ratio = hash_num("r" + key + task_id, 68, 96) / 100.0
    if is_revenue:
        target = hash_num(key, 20, 95) * 1000000000.0
        done = target * ratio
        return {
            "total": target,
            "done": done
        }
    done = int(base * ratio)
    return {
        "total": base,
        "done": done
    }


def get_task_prefix(applicable_teams):
    groups = set()
    for t in applicable_teams:
        if t.startswith("HKD"):
            groups.add("HKD")
        elif t.startswith("QLDN"):
            groups.add("QLDN")
        elif t == "KIEMTRA":
            groups.add("KT")
        elif t == "HCTH":
            groups.add("HC")
        elif t == "NVDTPC":
            groups.add("NV")
        elif t == "QLTK":
            groups.add("TK")
            
    if len(groups) >= 3:
        return "(Chung)"
    elif len(groups) == 1:
        return f"({list(groups)[0]})"
    elif len(groups) == 2:
        return f"({'-'.join(sorted(list(groups)))})"
    else:
        return "(Chung)"


def clean_sheet_name(task_title, prefix):
    # Sheet name limit is 31 characters
    clean_name = task_title
    for char in ['\\', '/', '?', '*', ':', '[', ']']:
        clean_name = clean_name.replace(char, '')
    max_len = 31 - len(prefix) - 1
    return f"{prefix} {clean_name[:max_len]}".strip()


def main():
    try:
        import inspect_monthly_data
    except Exception as e:
        print("Inspect warning:", e)
    print("Đang tạo file template mẫu...")
    create_template()

    # Load template
    wb = load_workbook(TEMPLATE_FILE)
    ws_dest = wb["Du lieu"]
    ws_dest.delete_rows(2, 305)  # Clear template rows

    print("Đang tạo các sheet nhiệm vụ nguồn chi tiết...")
    
    # Stylings
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    subtle_fill = PatternFill("solid", fgColor="F8FAFC")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    created_sheets = set()
    mst_counter = 8102345000
    cccd_counter = 1095012000
    
    # Store mapping of (task_title, officer_name) -> list of rows in task sheet
    task_sheet_mappings = []

    # Iterate over ALL tasks in seed_data.DEMO_TASKS
    for t_def in seed_data.DEMO_TASKS:
        task_code = t_def["Mã nhiệm vụ"]
        task_title = t_def["Tên nhiệm vụ"]
        has_detail = t_def.get("Có chi tiết", "Không") == "Có"
        
        # Find which teams this task applies to
        applicable_teams = []
        for team_id, task_codes in seed_data.DEMO_TASK_APP.items():
            if task_code in task_codes:
                applicable_teams.append(team_id)
                
        if not applicable_teams:
            continue
            
        # Determine prefix and sheet name
        prefix = get_task_prefix(applicable_teams)
        desired_name = clean_sheet_name(task_title, prefix)
        
        # Handle collisions
        sheet_name = desired_name[:31]
        if sheet_name in created_sheets:
            for idx in range(2, 100):
                suffix = f"_{idx}"
                candidate = desired_name[:31 - len(suffix)] + suffix
                if candidate not in created_sheets:
                    sheet_name = candidate
                    break
        created_sheets.add(sheet_name)
        
        # Create worksheet
        ws_task = wb.create_sheet(sheet_name)
        ws_task.sheet_view.showGridLines = True
        
        # Table Headers
        ws_task.append(["STT", "Cán bộ", "Tổ quản lý", "Địa bàn", "Mã số thuế", "CCCD", "Phải thực hiện", "Đã thực hiện", "Thời hạn", "Ghi chú"])
        for cell in ws_task[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        task_row_idx = 2
        for team_id in applicable_teams:
            staff_list = seed_data.DEMO_STAFF.get(team_id, [])
            for o_idx, s in enumerate(staff_list):
                is_revenue = (task_code == "THU")
                m = metric_for(team_id + s["name"], task_code, is_revenue, o_idx)
                
                assigned = round(m["total"])
                completed = round(m["done"])
                
                # Check if this task is detailed down to each taxpayer
                if has_detail and assigned > 0:
                    # Write multiple detailed taxpayer rows (e.g. 3 rows)
                    parts = min(3, assigned)
                    assigned_parts = [1] * parts
                    # Distribute remaining index if any
                    if assigned > parts:
                        assigned_parts[-1] = assigned - parts + 1
                        
                    completed_parts = [0] * parts
                    rem_completed = completed
                    for p_idx in range(parts):
                        take = min(assigned_parts[p_idx], rem_completed)
                        completed_parts[p_idx] = take
                        rem_completed -= take
                        
                    for p_idx in range(parts):
                        mst_counter += 1
                        cccd_counter += 1
                        mst_str = f"{mst_counter}"
                        cccd_str = f"00{cccd_counter}"
                        
                        ws_task.append([
                            task_row_idx - 1,
                            s["name"],
                            team_id,
                            s["area"],
                            mst_str,
                            cccd_str,
                            assigned_parts[p_idx],
                            completed_parts[p_idx],
                            "2026-12-31",
                            f"Người nộp thuế mẫu {p_idx + 1}"
                        ])
                        
                        task_sheet_mappings.append({
                            "task_title": task_title,
                            "sheet_name": sheet_name,
                            "row": task_row_idx
                        })
                        task_row_idx += 1
                else:
                    # Summary row (no taxpayer details)
                    ws_task.append([
                        task_row_idx - 1,
                        s["name"],
                        team_id,
                        s["area"],
                        "",
                        "",
                        assigned,
                        completed,
                        "2026-12-31",
                        ""
                    ])
                    
                    task_sheet_mappings.append({
                        "task_title": task_title,
                        "sheet_name": sheet_name,
                        "row": task_row_idx
                    })
                    task_row_idx += 1
                    
        # Format the task sheet rows
        style_range(ws_task, task_row_idx - 1, 10, border)
        for row in range(2, task_row_idx):
            ws_task[f"A{row}"].fill = subtle_fill
            ws_task[f"A{row}"].font = Font(color="475569", bold=True)
            ws_task[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
            ws_task[f"G{row}"].number_format = "#,##0"
            ws_task[f"H{row}"].number_format = "#,##0"
            ws_task[f"I{row}"].number_format = "yyyy-mm-dd"
            
        # Column widths
        for col, width in {"A": 8, "B": 24, "C": 18, "D": 24, "E": 18, "F": 18, "G": 16, "H": 16, "I": 16, "J": 28}.items():
            ws_task[col + "1"].alignment = Alignment(horizontal="center")
            ws_task.column_dimensions[col].width = width

    # 2. Populate main "Du lieu" sheet with formulas pointing to task sheets
    print("Đang tạo liên kết công thức Excel cho sheet Du lieu đầu tiên...")
    row_idx = 2
    for mapping in task_sheet_mappings:
        task_title = mapping["task_title"]
        sh_name = mapping["sheet_name"]
        o_row = mapping["row"]
        
        # Formulas linking to the task sheet
        officer_formula = f"='{sh_name}'!B{o_row}"
        team_formula = f"='{sh_name}'!C{o_row}"
        new_area_formula = f"='{sh_name}'!D{o_row}"
        mst_formula = f"='{sh_name}'!E{o_row}"
        cccd_formula = f"='{sh_name}'!F{o_row}"
        assigned_formula = f"='{sh_name}'!G{o_row}"
        completed_formula = f"='{sh_name}'!H{o_row}"
        deadline_formula = f"='{sh_name}'!I{o_row}"
        note_formula = f"='{sh_name}'!J{o_row}"

        ws_dest.append([
            row_idx - 1,          # STT
            task_title,           # Tên nhiệm vụ
            officer_formula,      # Cán bộ
            "",                   # Địa bàn xã cũ
            new_area_formula,     # Địa bàn xã mới
            team_formula,         # Tổ quản lý
            mst_formula,          # Mã số thuế
            cccd_formula,         # CCCD
            assigned_formula,     # Phải thực hiện
            completed_formula,    # Đã thực hiện
            deadline_formula,     # Thời hạn
            note_formula          # Ghi chú
        ])
        row_idx += 1

    # Style and format the main sheet rows
    style_range(ws_dest, row_idx - 1, len(HEADERS), border)
    for row in range(2, row_idx):
        ws_dest[f"A{row}"].fill = subtle_fill
        ws_dest[f"A{row}"].font = Font(color="475569", bold=True)
        ws_dest[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws_dest[f"I{row}"].number_format = "#,##0"
        ws_dest[f"J{row}"].number_format = "#,##0"
        ws_dest[f"K{row}"].number_format = "yyyy-mm-dd"

    # Add back the table styling
    from openpyxl.worksheet.table import Table, TableStyleInfo
    table = Table(displayName="BangTongHopDayDuLieu", ref=f"A1:L{row_idx - 1}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws_dest.add_table(table)

    # 3. Create "Bao cao tuan" sheet
    print("Đang tạo sheet Báo cáo tuần tổng hợp...")
    ws_report = wb.create_sheet("Bao cao tuan", 1)  # Insert as second sheet (index 1)
    ws_report.sheet_view.showGridLines = True
    
    # Title
    ws_report.merge_cells("A1:G1")
    title_cell = ws_report["A1"]
    title_cell.value = "BÁO CÁO TUẦN - ĐÁNH GIÁ CHỈ TIÊU CÔNG VIỆC THUẾ CƠ SỞ 23"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_report.row_dimensions[1].height = 40
    
    # Headers
    ws_report.append(["STT", "Tên nhiệm vụ", "Tổ thực hiện", "Tổng chỉ tiêu giao", "Tổng thực hiện", "Tỷ lệ hoàn thành", "Đánh giá tiến độ"])
    for cell in ws_report[2]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws_report.row_dimensions[2].height = 28
    
    # Rows for Bao cao tuan
    WEEKLY_CRITERIA = [
        {"stt": "1.0", "title": "Tỷ lệ số thực thu NSNN", "task_title": "Theo dõi số thu ngân sách"},
        {"stt": "4.0", "title": "Tỷ lệ nợ tăng/giảm so với đầu năm (biên độ chênh lệch 20%)", "task_title": "Quản lý nợ thuế"},
        {"stt": "9.0", "title": "Tỷ lệ số lượng NNT đã thực hiện cưỡng chế/Số lượng phải thực hiện cưỡng chế", "task_title": "Cưỡng chế nợ thuế"},
        {"stt": "10.0", "title": "Tỷ lệ số tiền đã thực hiện cưỡng chế/Số tiền phải thực hiện cưỡng chế", "task_title": "Cưỡng chế nợ thuế"},
        {"stt": "12.0", "title": "Tỷ lệ số tiền đã thực hiện tạm hoãn xuất cảnh/Số tiền phải thực hiện THXC", "task_title": "Tạm hoãn xuất cảnh"},
        {"stt": "15.0", "title": "Tỷ lệ tồn quá hạn thực hiện TTHC (Đã loại trừ hoàn GTGT và TNCN, đóng mã)", "task_title": "Giải quyết thủ tục hành chính"},
        {"stt": "16.0", "title": "Tỷ lệ đánh giá sự hài lòng (hệ thống Cục Thuế)", "task_title": "Theo dõi sự hài lòng"},
        {"stt": "17.0", "title": "Rà soát TPR", "task_title": "Rà soát TPR"},
        {"stt": "18.0", "title": "Xác minh hóa đơn", "task_title": "Rà soát, xác minh và xử lý hóa đơn / Xác minh hóa đơn"},
        {"stt": "19.0", "title": "Hệ số K", "task_title": "Xử lý cảnh báo hệ số K"},
        {"stt": "22.0", "title": "Đóng mã giải thể phát sinh mới", "task_title": "Đăng ký thuế, trạng thái MST và đóng mã / Giải thể"},
        {"stt": "21.0", "title": "Đóng mã giải thể theo gói của Phòng QLDN3", "task_title": "Chuyên đề và gói dữ liệu doanh nghiệp"},
        {"stt": "25.0", "title": "Tỷ lệ số tồn quá hạn thuế TNCN", "task_title": "Hoàn thuế TNCN"},
        {"stt": "23.0", "title": "Tỷ lệ hoàn thành gói TNCN nhiều nguồn 2025", "task_title": "Xử lý gói nhiều nguồn 2025"},
        {"stt": "34.0", "title": "Tỷ lệ hoàn thành chỉ tiêu kiểm tra tại trụ sở NNT", "task_title": "Kiểm tra tại trụ sở cơ quan thuế theo kế hoạch / Kiểm tra tại trụ sở cơ quan thuế"},
        {"stt": "35.0", "title": "Tỷ lệ đôn đốc nộp truy thu, xử phạt", "task_title": "Xử phạt VPHC"},
    ]

    report_row_idx = 3
    for crit in WEEKLY_CRITERIA:
        task_title = crit["task_title"]
        crit_title = crit["title"]
        crit_stt = crit["stt"]
        
        actual_sh_name = None
        for m in task_sheet_mappings:
            if m["task_title"] == task_title:
                actual_sh_name = m["sheet_name"]
                break
        if not actual_sh_name:
            continue
            
        task_ws = wb[actual_sh_name]
        last_row = task_ws.max_row
        
        # Formulas referencing the task sheet
        assigned_formula = f"=SUM('{actual_sh_name}'!G2:G{last_row})"
        completed_formula = f"=SUM('{actual_sh_name}'!H2:H{last_row})"
        rate_formula = f"=IF(D{report_row_idx}>0, E{report_row_idx}/D{report_row_idx}, 0)"
        status_formula = f'=IF(F{report_row_idx}>=0.9, "Đạt (>=90%)", "Chưa đạt (<90%)")'
        
        # Determine responsible teams
        applicable_teams = []
        for team_id, task_codes in seed_data.DEMO_TASK_APP.items():
            task_code = None
            for t_def in seed_data.DEMO_TASKS:
                if t_def["Tên nhiệm vụ"] == task_title:
                    task_code = t_def["Mã nhiệm vụ"]
                    break
            if task_code in task_codes:
                applicable_teams.append(team_id)
                
        ws_report.append([
            crit_stt,
            crit_title,
            ", ".join(applicable_teams),
            assigned_formula,
            completed_formula,
            rate_formula,
            status_formula
        ])
        
        style_range(ws_report, report_row_idx, 7, border)
        ws_report[f"A{report_row_idx}"].fill = subtle_fill
        ws_report[f"A{report_row_idx}"].font = Font(color="475569", bold=True)
        ws_report[f"A{report_row_idx}"].alignment = Alignment(horizontal="center")
        ws_report[f"D{report_row_idx}"].number_format = "#,##0"
        ws_report[f"E{report_row_idx}"].number_format = "#,##0"
        ws_report[f"F{report_row_idx}"].number_format = "0.0%"
        ws_report[f"G{report_row_idx}"].alignment = Alignment(horizontal="center")
        
        report_row_idx += 1
        
    for col, width in {"A": 8, "B": 32, "C": 24, "D": 20, "E": 20, "F": 20, "G": 20}.items():
        ws_report.column_dimensions[col].width = width
        
    table_rep = Table(displayName="BangBaoCaoTuan", ref=f"A2:G{report_row_idx - 1}")
    table_rep.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws_report.add_table(table_rep)

    # 4. Create "Bao cao KPI tuan" sheet
    print("Đang tạo sheet Báo cáo KPI tuần theo thiết kế của văn phòng...")
    ws_kpi = wb.create_sheet("Bao cao KPI tuan", 2)
    ws_kpi.sheet_view.showGridLines = True

    kpi_headers = [
        {"ref": "A1:A2", "val": "Đơn vị"},
        {"ref": "B1:D1", "val": "Chỉ tiêu Tỷ lệ số thực thu NSNN"},
        {"ref": "E1:E2", "val": "Tổng điểm trung bình thu NSNN"},
        {"ref": "F1:H1", "val": "Tỷ lệ nợ tăng/giảm so với đầu năm (biên độ chênh lệch 20%)"},
        {"ref": "I1:K1", "val": "Tỷ lệ số lượng NNT đã thực hiện cưỡng chế/Số lượng phải thực hiện cưỡng chế"},
        {"ref": "L1:N1", "val": "Tỷ lệ số tiền đã thực hiện cưỡng chế/Số tiền phải thực hiện cưỡng chế"},
        {"ref": "O1:Q1", "val": "Tỷ lệ số tiền đã thực hiện tạm hoãn xuất cảnh/Số tiền phải thực hiện THXC"},
        {"ref": "R1:R2", "val": "Tổng điểm trung bình công tác QLN"},
        {"ref": "S1:U1", "val": "Tỷ lệ tồn quá hạn thực hiện TTHC (Đã loại trừ hoàn GTGT và TNCN, đóng mã) "},
        {"ref": "V1:X1", "val": "Tỷ lệ đánh giá sự hài lòng (hệ thống Cục Thuế)"},
        {"ref": "Y1:Y2", "val": "Tổng điểm trung bình TTHC"},
        {"ref": "Z1:AB1", "val": "Rà soát TPR"},
        {"ref": "AC1:AE1", "val": "Xác minh hóa đơn"},
        {"ref": "AF1:AH1", "val": "Hệ số K"},
        {"ref": "AI1:AK1", "val": "Đóng mã giải thể phát sinh mới"},
        {"ref": "AL1:AN1", "val": "Đóng mã giải thể theo gói của Phòng QLDN3"},
        {"ref": "AO1:AQ1", "val": "Tỷ lệ số tồn quá hạn thuế TNCN"},
        {"ref": "AR1:AT1", "val": "Tỷ lệ hoàn thành gói TNCN nhiều nguồn 2025"},
        {"ref": "AU1:AU2", "val": "Tổng điểm trung bình công tác quản lý hỗ trợ"},
        {"ref": "AV1:AX1", "val": "Tỷ lệ hoàn thành chỉ tiêu kiểm tra tại trụ sở NNT"},
        {"ref": "AY1:BA1", "val": "Tỷ lệ đôn đốc nộp truy thu, xử phạt"},
        {"ref": "BB1:BB2", "val": "Tổng điểm trung bình công tác kiểm tra"},
        {"ref": "BC1:BC2", "val": "Tổng điểm"},
        {"ref": "BD1:BD2", "val": "Xếp hạng"},
    ]

    subheaders = {
        "B": "Tỷ lệ số thực thu NSNN", "C": "Bình quân toàn ngành", "D": "Tính điểm",
        "F": "Tỷ lệ nợ tăng/giảm so với đầu năm", "G": "Bình quân toàn ngành", "H": "Tính điểm",
        "I": "Tỷ lệ số lượng NNT", "J": "Bình quân toàn ngành", "K": "Tính điểm",
        "L": "Tỷ lệ số tiền", "M": "Bình quân toàn ngành", "N": "Tính điểm",
        "O": "Tỷ lệ số tiền", "P": "Bình quân toàn ngành", "Q": "Tính điểm",
        "S": "Tỷ lệ tồn quá hạn", "T": "Bình quân toàn ngành", "U": "Tính điểm",
        "V": "Tỷ lệ hài lòng", "W": "So sánh với 90%", "X": "Tính điểm",
        "Z": "Rà soát TPR", "AA": "Bình quân toàn ngành", "AB": "Tính điểm",
        "AC": "Xác minh hóa đơn", "AD": "Bình quân toàn ngành", "AE": "Tính điểm",
        "AF": "Hệ số K", "AG": "Bình quân toàn ngành", "AH": "Tính điểm",
        "AI": "Tỷ lệ hồ sơ đã đóng mã", "AJ": "Bình quân toàn ngành", "AK": "Tính điểm",
        "AL": "Tỷ lệ hoàn thành", "AM": "Bình quân toàn ngành", "AN": "Tính điểm",
        "AO": "Tỷ lệ số tồn quá hạn", "AP": "Bình quân toàn ngành", "AQ": "Tính điểm",
        "AR": "Tỷ lệ hoàn thành", "AS": "Bình quân toàn ngành", "AT": "Tính điểm",
        "AV": "Tỷ lệ hoàn thành", "AW": "Bình quân toàn ngành", "AX": "Tính điểm",
        "AY": "Tỷ lệ đôn đốc", "AZ": "Bình quân toàn ngành", "BA": "Tính điểm",
    }

    # Style definitions
    kpi_header_fill = PatternFill("solid", fgColor="1D4ED8")
    kpi_sub_fill = PatternFill("solid", fgColor="DBEAFE")
    kpi_avg_fill = PatternFill("solid", fgColor="F1F5F9")
    kpi_score_fill = PatternFill("solid", fgColor="FEF08A")
    kpi_green_fill = PatternFill("solid", fgColor="BBF7D0")
    kpi_header_font = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
    kpi_sub_font = Font(name="Calibri", size=9, color="1E3A8A", bold=True)
    kpi_body_font = Font(name="Calibri", size=10)
    kpi_bold_font = Font(name="Calibri", size=10, bold=True)
    kpi_border = Border(top=thin, left=thin, right=thin, bottom=thin)

    def merge_and_format(ws, ref, val, fill, font, alignment):
        ws.merge_cells(ref)
        first_cell = ws[ref.split(":")[0]]
        first_cell.value = val
        for r_cells in ws[ref]:
            for cell in r_cells:
                cell.fill = fill
                cell.font = font
                cell.alignment = alignment
                cell.border = kpi_border

    # Format row dimensions
    ws_kpi.row_dimensions[1].height = 28
    ws_kpi.row_dimensions[2].height = 28

    # Apply merges and formats for Row 1
    for h in kpi_headers:
        ref = h["ref"]
        val = h["val"]
        is_green = val.startswith("Tổng điểm") or val in ["Tổng điểm", "Xếp hạng"]
        fill = kpi_green_fill if is_green else kpi_header_fill
        font = Font(name="Calibri", size=10, color="065F46" if is_green else "FFFFFF", bold=True)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        merge_and_format(ws_kpi, ref, val, fill, font, align)

    # Write subheaders in Row 2
    for col_letter, sub_val in subheaders.items():
        cell = ws_kpi[f"{col_letter}2"]
        cell.value = sub_val
        is_score = sub_val == "Tính điểm"
        cell.fill = kpi_score_fill if is_score else kpi_sub_fill
        cell.font = kpi_sub_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = kpi_border

    # Mapping of indicators to task titles
    indicators_map = [
        {"rate_col": "B", "bq_val": 0.71, "type": "std_revenue", "task": "Theo dõi số thu ngân sách"},
        {"rate_col": "F", "bq_val": 0.17, "type": "debt", "task": "Quản lý nợ thuế"},
        {"rate_col": "I", "bq_val": 0.56, "type": "std", "task": "Cưỡng chế nợ thuế"},
        {"rate_col": "L", "bq_val": 0.90, "type": "std", "task": "Cưỡng chế nợ thuế"},
        {"rate_col": "O", "bq_val": 0.837, "type": "std", "task": "Tạm hoãn xuất cảnh"},
        {"rate_col": "S", "bq_val": 0.10558, "type": "rev", "task": "Giải quyết thủ tục hành chính"},
        {"rate_col": "V", "bq_val": 0.90, "type": "sat", "task": "Theo dõi sự hài lòng"},
        {"rate_col": "Z", "bq_val": 0.994, "type": "std", "task": "Rà soát TPR"},
        {"rate_col": "AC", "bq_val": 0.9923, "type": "std", "task": "Rà soát, xác minh và xử lý hóa đơn / Xác minh hóa đơn"},
        {"rate_col": "AF", "bq_val": 0.9916, "type": "std", "task": "Xử lý cảnh báo hệ số K"},
        {"rate_col": "AI", "bq_val": 0.50, "type": "std", "task": "Đăng ký thuế, trạng thái MST và đóng mã / Giải thể"},
        {"rate_col": "AL", "bq_val": 0.2909, "type": "std", "task": "Chuyên đề và gói dữ liệu doanh nghiệp"},
        {"rate_col": "AO", "bq_val": 0.0075, "type": "rev", "task": "Hoàn thuế TNCN"},
        {"rate_col": "AR", "bq_val": 1.0, "type": "std", "task": "Xử lý gói nhiều nguồn 2025"},
        {"rate_col": "AV", "bq_val": 0.62, "type": "std", "task": "Kiểm tra tại trụ sở cơ quan thuế theo kế hoạch / Kiểm tra tại trụ sở cơ quan thuế"},
        {"rate_col": "AY", "bq_val": 0.92, "type": "std", "task": "Xử phạt VPHC"},
    ]

    def get_sh_name(task_title):
        for m in task_sheet_mappings:
            if m["task_title"] == task_title:
                return m["sheet_name"]
        return ""

    teams_list = [t["code"] for t in seed_data.DEMO_TEAMS]
    
    # Write rows for each team
    for idx, team_id in enumerate(teams_list):
        row_num = 3 + idx
        team_name = [t["name"] for t in seed_data.DEMO_TEAMS if t["code"] == team_id][0]
        
        # Write team name (Col A)
        cell_name = ws_kpi.cell(row_num, 1, team_name)
        cell_name.font = kpi_bold_font
        cell_name.border = kpi_border
        cell_name.fill = subtle_fill
        ws_kpi.row_dimensions[row_num].height = 24
        
        # Write indicators
        for ind in indicators_map:
            r_col = ind["rate_col"]
            idx_num = column_index_from_string(r_col)
            bq_col = get_column_letter(idx_num + 1)
            sc_col = get_column_letter(idx_num + 2)
            
            sh_name = get_sh_name(ind["task"])
            if sh_name:
                task_ws = wb[sh_name]
                last_row = task_ws.max_row
                # Formula to calculate team average rate
                rate_formula = f"=IF(SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!G2:G{last_row})>0, SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!H2:H{last_row}) / SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!G2:G{last_row}), 0)"
            else:
                rate_formula = 0
                
            # 1. Rate Cell
            c_rate = ws_kpi[f"{r_col}{row_num}"]
            c_rate.value = rate_formula
            c_rate.number_format = "0.0%"
            c_rate.font = kpi_body_font
            c_rate.border = kpi_border
            c_rate.alignment = Alignment(horizontal="right")
            
            # 2. Benchmark Cell
            c_bq = ws_kpi[f"{bq_col}{row_num}"]
            c_bq.value = ind["bq_val"]
            c_bq.number_format = "0.0%" if ind["type"] != "sat" else "0.0%"
            c_bq.font = kpi_body_font
            c_bq.border = kpi_border
            c_bq.fill = kpi_avg_fill
            c_bq.alignment = Alignment(horizontal="right")
            
            # 3. Score Cell
            c_score = ws_kpi[f"{sc_col}{row_num}"]
            if ind["type"] == "std_revenue":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, 1, IF({r_col}{row_num}>={bq_col}{row_num}, 0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, 0, IF({r_col}{row_num}>={bq_col}{row_num}-0.1, -0.5, -1))))"
            elif ind["type"] == "debt":
                score_formula = f"=IF({r_col}{row_num}>=0.2, -1, IF({r_col}{row_num}>=0, -0.5, IF({r_col}{row_num}>=-0.2, 0.5, 1)))"
            elif ind["type"] == "std":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, 1, IF({r_col}{row_num}>={bq_col}{row_num}, 0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, -0.5, -1)))"
            elif ind["type"] == "rev":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, -1, IF({r_col}{row_num}>={bq_col}{row_num}, -0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, 0.5, 1)))"
            elif ind["type"] == "sat":
                score_formula = f"=IF({r_col}{row_num}>=0.95, 1, IF({r_col}{row_num}>=0.9, 0.5, IF({r_col}{row_num}>=0.85, -0.5, -1)))"
            else:
                score_formula = 0
                
            c_score.value = score_formula
            c_score.number_format = "0.0"
            c_score.font = kpi_bold_font
            c_score.border = kpi_border
            c_score.fill = kpi_score_fill
            c_score.alignment = Alignment(horizontal="right")

        # Write Group Averages / Total Averages Formulas
        # E: =D
        ws_kpi[f"E{row_num}"] = f"=D{row_num}"
        # R: =AVERAGE(H, K, N, Q)
        ws_kpi[f"R{row_num}"] = f"=AVERAGE(H{row_num}, K{row_num}, N{row_num}, Q{row_num})"
        # Y: =AVERAGE(U, X)
        ws_kpi[f"Y{row_num}"] = f"=AVERAGE(U{row_num}, X{row_num})"
        # AU: =AVERAGE(AB, AE, AH, AK, AN, AQ, AT)
        ws_kpi[f"AU{row_num}"] = f"=AVERAGE(AB{row_num}, AE{row_num}, AH{row_num}, AK{row_num}, AN{row_num}, AQ{row_num}, AT{row_num})"
        # BB: =AVERAGE(AX, BA)
        ws_kpi[f"BB{row_num}"] = f"=AVERAGE(AX{row_num}, BA{row_num})"
        # BC (Total Score): =SUM(E, R, Y, AU, BB)
        ws_kpi[f"BC{row_num}"] = f"=SUM(E{row_num}, R{row_num}, Y{row_num}, AU{row_num}, BB{row_num})"
        # BD (Ranking): =RANK(BC, BC$3:BC$10)
        ws_kpi[f"BD{row_num}"] = f"=RANK(BC{row_num}, BC$3:BC$10)"

        # Style Group Average Columns
        for g_col in ["E", "R", "Y", "AU", "BB", "BC", "BD"]:
            cell = ws_kpi[f"{g_col}{row_num}"]
            cell.fill = kpi_green_fill
            cell.font = kpi_bold_font
            cell.border = kpi_border
            cell.number_format = "0.00" if g_col != "BD" else "0"
            cell.alignment = Alignment(horizontal="right")

    # Column widths for weekly KPI
    ws_kpi.column_dimensions["A"].width = 30
    for c_idx in range(2, 57):
        col_letter = get_column_letter(c_idx)
        ws_kpi.column_dimensions[col_letter].width = 11

    # 5. Create "Bao cao KPI thang" sheet
    print("Đang tạo sheet Báo cáo KPI tháng căn cứ theo Bộ tiêu chí và CV Chỉ tiêu TCS23...")
    ws_kpi_month = wb.create_sheet("Bao cao KPI thang", 3)
    ws_kpi_month.sheet_view.showGridLines = True

    kpi_month_headers = [
        {"ref": "A1:A2", "val": "Đơn vị (Tổ)"},
        {"ref": "B1:D1", "val": "1. Tỷ lệ số thực thu NSNN (Lũy kế/DTPL)"},
        {"ref": "E1:H1", "val": "2. Tỷ lệ thực thu tháng/ước thu tháng (chênh lệch 3%)"},
        {"ref": "I1:K1", "val": "3. Tỷ lệ thu NSNN tăng giảm so với cùng kỳ"},
        {"ref": "L1:L2", "val": "Tổng điểm TB Thu NSNN"},
        {"ref": "M1:O1", "val": "4. Tỷ lệ nợ tăng/giảm so với 31/12/2025"},
        {"ref": "P1:R1", "val": "5. Tăng/giảm nợ so với tháng trước"},
        {"ref": "S1:U1", "val": "6. Tỷ lệ Tổng nợ / Tổng thu NSNN"},
        {"ref": "V1:X1", "val": "7. Tỷ lệ Nợ khả năng thu / Tổng thu NSNN"},
        {"ref": "Y1:AA1", "val": "8. Tỷ lệ nợ Thuế, phí / Tổng thu NSNN (<5%)"},
        {"ref": "AB1:AD1", "val": "9. Tỷ lệ số lượng NNT cưỡng chế"},
        {"ref": "AE1:AG1", "val": "10. Tỷ lệ số tiền cưỡng chế"},
        {"ref": "AH1:AJ1", "val": "11. Tỷ lệ số lượng tạm hoãn xuất cảnh (TT00, 03, 06)"},
        {"ref": "AK1:AM1", "val": "12. Tỷ lệ số tiền tạm hoãn xuất cảnh (TT00, 03, 06)"},
        {"ref": "AN1:AN2", "val": "Tổng điểm TB Quản lý nợ"},
        {"ref": "AO1:AQ1", "val": "14. Tỷ lệ tồn quá hạn TTHC theo tháng"},
        {"ref": "AR1:AT1", "val": "16. Tỷ lệ đánh giá sự hài lòng (Hệ thống Cục Thuế)"},
        {"ref": "AU1:AW1", "val": "Xử lý phản ánh DVCQG & Hỏi đáp BTC"},
        {"ref": "AX1:AX2", "val": "Tổng điểm TB TTHC"},
        {"ref": "AY1:BA1", "val": "17. Rà soát TPR"},
        {"ref": "BB1:BD1", "val": "18. Xác minh hóa đơn"},
        {"ref": "BE1:BG1", "val": "19. Hệ số K"},
        {"ref": "BH1:BJ1", "val": "20. Xử phạt VPHC"},
        {"ref": "BK1:BM1", "val": "21. Đóng mã giải thể theo gói QLDN3"},
        {"ref": "BN1:BP1", "val": "22. Đóng mã giải thể mới"},
        {"ref": "BQ1:BS1", "val": "23. Gói nhiều nguồn 2025"},
        {"ref": "BT1:BV1", "val": "24. Kiểm tra tại bàn"},
        {"ref": "BW1:BY1", "val": "Hoàn thuế TNCN tự động & Hoàn GTGT điện tử"},
        {"ref": "BZ1:BZ2", "val": "Tổng điểm TB Quản lý DN"},
        {"ref": "CA1:CC1", "val": "27. Tổng thu HKD so với DTPL"},
        {"ref": "CD1:CF1", "val": "28. Kê khai thuế HKD đúng hạn"},
        {"ref": "CG1:CI1", "val": "29. Đăng ký & sử dụng eTax Mobile HKD"},
        {"ref": "CJ1:CL1", "val": "30. HĐĐT máy tính tiền HKD"},
        {"ref": "CM1:CO1", "val": "31. Khai thác TMĐT các gói trọng điểm"},
        {"ref": "CP1:CR1", "val": "32. Chuyển đổi HKD lên DN"},
        {"ref": "CS1:CS2", "val": "Tổng điểm TB Hộ kinh doanh"},
        {"ref": "CT1:CV1", "val": "34. Hoàn thành kế hoạch kiểm tra tại trụ sở"},
        {"ref": "CW1:CY1", "val": "Kế hoạch kiểm tra đích danh 2026"},
        {"ref": "CZ1:DB1", "val": "35. Đôn đốc nộp truy thu, xử phạt sau KT"},
        {"ref": "DC1:DE1", "val": "37. Kiểm tra đóng mã giải thể"},
        {"ref": "DF1:DF2", "val": "Tổng điểm TB Kiểm tra"},
        {"ref": "DG1:DG2", "val": "Tổng điểm thi đua KPI tháng"},
        {"ref": "DH1:DH2", "val": "Xếp hạng tháng"},
    ]

    subheaders_month = {
        "B": "Tỷ lệ", "C": "BQ ngành", "D": "Tính điểm",
        "E": "Thực hiện", "F": "Ước thu", "G": "Chênh lệch", "H": "Tính điểm",
        "I": "Tỷ lệ", "J": "BQ ngành", "K": "Tính điểm",
        "M": "Tỷ lệ", "N": "BQ ngành", "O": "Tính điểm",
        "P": "Tỷ lệ", "Q": "BQ ngành", "R": "Tính điểm",
        "S": "Tỷ lệ", "T": "Cục giao (8%)", "U": "Tính điểm",
        "V": "Tỷ lệ", "W": "Cục giao (5%)", "X": "Tính điểm",
        "Y": "Tỷ lệ", "Z": "Cục giao (5%)", "AA": "Tính điểm",
        "AB": "Tỷ lệ SL", "AC": "BQ ngành (56%)", "AD": "Tính điểm",
        "AE": "Tỷ lệ Tiền", "AF": "BQ ngành (90%)", "AG": "Tính điểm",
        "AH": "Tỷ lệ SL", "AI": "BQ ngành (80%)", "AJ": "Tính điểm",
        "AK": "Tỷ lệ Tiền", "AL": "BQ ngành (83.7%)", "AM": "Tính điểm",
        "AO": "Tỷ lệ tồn", "AP": "BQ ngành (9.8%)", "AQ": "Tính điểm",
        "AR": "Tỷ lệ hài lòng", "AS": "Chuẩn 90%", "AT": "Tính điểm",
        "AU": "Tỷ lệ đúng hạn", "AV": "Chuẩn 100%", "AW": "Tính điểm",
        "AY": "Tỷ lệ", "AZ": "BQ ngành (99.4%)", "BA": "Tính điểm",
        "BB": "Tỷ lệ", "BC": "BQ ngành (99.4%)", "BD": "Tính điểm",
        "BE": "Tỷ lệ", "BF": "BQ ngành (99.4%)", "BG": "Tính điểm",
        "BH": "Tỷ lệ", "BI": "BQ ngành (79.7%)", "BJ": "Tính điểm",
        "BK": "Tỷ lệ", "BL": "BQ ngành (49.1%)", "BM": "Tính điểm",
        "BN": "Tỷ lệ", "BO": "BQ ngành (62.8%)", "BP": "Tính điểm",
        "BQ": "Tỷ lệ", "BR": "Chuẩn 100%", "BS": "Tính điểm",
        "BT": "Tỷ lệ", "BU": "BQ ngành (20.7%)", "BV": "Tính điểm",
        "BW": "Tỷ lệ", "BX": "Chuẩn 95%", "BY": "Tính điểm",
        "CA": "Tỷ lệ", "CB": "Chuẩn 100%", "CC": "Tính điểm",
        "CD": "Tỷ lệ", "CE": "BQ ngành (97.3%)", "CF": "Tính điểm",
        "CG": "Tỷ lệ", "CH": "BQ ngành (88.8%)", "CI": "Tính điểm",
        "CJ": "Tỷ lệ", "CK": "BQ ngành (87.0%)", "CL": "Tính điểm",
        "CM": "Tỷ lệ", "CN": "Chuẩn 85%", "CO": "Tính điểm",
        "CP": "Tỷ lệ", "CQ": "Chuẩn 100%", "CR": "Tính điểm",
        "CT": "Tỷ lệ", "CU": "BQ ngành (83.1%)", "CV": "Tính điểm",
        "CW": "Tỷ lệ", "CX": "BQ ngành (39.8%)", "CY": "Tính điểm",
        "CZ": "Tỷ lệ", "DA": "BQ ngành (94.7%)", "DB": "Tính điểm",
        "DC": "Tỷ lệ", "DD": "BQ ngành (82.0%)", "DE": "Tính điểm",
    }

    ws_kpi_month.row_dimensions[1].height = 28
    ws_kpi_month.row_dimensions[2].height = 28

    for h in kpi_month_headers:
        ref = h["ref"]
        val = h["val"]
        is_green = val.startswith("Tổng điểm") or val in ["Tổng điểm thi đua KPI tháng", "Xếp hạng tháng"]
        fill = kpi_green_fill if is_green else kpi_header_fill
        font = Font(name="Calibri", size=10, color="065F46" if is_green else "FFFFFF", bold=True)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        merge_and_format(ws_kpi_month, ref, val, fill, font, align)

    for col_letter, sub_val in subheaders_month.items():
        cell = ws_kpi_month[f"{col_letter}2"]
        cell.value = sub_val
        is_score = sub_val == "Tính điểm"
        cell.fill = kpi_score_fill if is_score else kpi_sub_fill
        cell.font = kpi_sub_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = kpi_border

    month_indicators_map = [
        {"rate_col": "B", "bq_val": 0.7964, "type": "std_revenue", "task": "Theo dõi số thu ngân sách"},
        {"rate_col": "I", "bq_val": 1.0771, "type": "std_revenue", "task": "Theo dõi số thu ngân sách"},
        {"rate_col": "M", "bq_val": 0.0456, "type": "debt_growth", "task": "Quản lý nợ thuế"},
        {"rate_col": "P", "bq_val": -0.0982, "type": "debt_growth", "task": "Tăng giảm nợ so với tháng trước và thu nợ theo kế hoạch"},
        {"rate_col": "S", "bq_val": 0.08, "type": "debt_ratio", "task": "Quản lý nợ thuế"},
        {"rate_col": "V", "bq_val": 0.05, "type": "debt_ratio", "task": "Quản lý nợ thuế"},
        {"rate_col": "Y", "bq_val": 0.05, "type": "debt_ratio", "task": "Quản lý nợ thuế"},
        {"rate_col": "AB", "bq_val": 0.56, "type": "std", "task": "Cưỡng chế nợ thuế"},
        {"rate_col": "AE", "bq_val": 0.90, "type": "std", "task": "Cưỡng chế nợ thuế"},
        {"rate_col": "AH", "bq_val": 0.80, "type": "std", "task": "Tạm hoãn xuất cảnh theo trạng thái MST (TT00, 03, 06)"},
        {"rate_col": "AK", "bq_val": 0.837, "type": "std", "task": "Tạm hoãn xuất cảnh"},
        {"rate_col": "AO", "bq_val": 0.098, "type": "overdue", "task": "Giải quyết thủ tục hành chính"},
        {"rate_col": "AR", "bq_val": 0.90, "type": "sat", "task": "Theo dõi sự hài lòng"},
        {"rate_col": "AU", "bq_val": 1.00, "type": "std", "task": "Xử lý phản ánh Cổng DVCQG và hỏi đáp Cổng TTĐT BTC"},
        {"rate_col": "AY", "bq_val": 0.9937, "type": "std", "task": "Rà soát TPR"},
        {"rate_col": "BB", "bq_val": 0.9938, "type": "std", "task": "Rà soát, xác minh và xử lý hóa đơn / Xác minh hóa đơn"},
        {"rate_col": "BE", "bq_val": 0.9943, "type": "std", "task": "Xử lý cảnh báo hệ số K"},
        {"rate_col": "BH", "bq_val": 0.7973, "type": "std", "task": "Xử phạt VPHC"},
        {"rate_col": "BK", "bq_val": 0.4908, "type": "std", "task": "Chuyên đề và gói dữ liệu doanh nghiệp"},
        {"rate_col": "BN", "bq_val": 0.628, "type": "std", "task": "Đăng ký thuế, trạng thái MST và đóng mã / Giải thể"},
        {"rate_col": "BQ", "bq_val": 1.00, "type": "std", "task": "Xử lý gói nhiều nguồn 2025"},
        {"rate_col": "BT", "bq_val": 0.2075, "type": "std", "task": "Kiểm tra tại trụ sở cơ quan thuế theo kế hoạch / Kiểm tra tại trụ sở cơ quan thuế"},
        {"rate_col": "BW", "bq_val": 0.95, "type": "std", "task": "Giải quyết hoàn thuế TNCN tự động"},
        {"rate_col": "CA", "bq_val": 1.00, "type": "std_revenue", "task": "Theo dõi số thu ngân sách"},
        {"rate_col": "CD", "bq_val": 0.973, "type": "std", "task": "Quản lý kê khai, tờ khai lỗi và xử phạt"},
        {"rate_col": "CG", "bq_val": 0.888, "type": "std", "task": "eTax Mobile"},
        {"rate_col": "CJ", "bq_val": 0.870, "type": "std", "task": "Hóa đơn điện tử có mã từ máy tính tiền của HKD"},
        {"rate_col": "CM", "bq_val": 0.85, "type": "std", "task": "Khai thác dữ liệu TMĐT các gói trọng điểm (GHTK, TikTok, CV 3049, Gói 3 1577)"},
        {"rate_col": "CP", "bq_val": 1.00, "type": "std", "task": "Tuyên truyền vận động hộ kinh doanh chuyển đổi lên doanh nghiệp"},
        {"rate_col": "CT", "bq_val": 0.8307, "type": "std", "task": "Thực hiện kế hoạch kiểm tra"},
        {"rate_col": "CW", "bq_val": 0.3975, "type": "std", "task": "Kế hoạch kiểm tra đích danh tại trụ sở NNT"},
        {"rate_col": "CZ", "bq_val": 0.9473, "type": "std", "task": "Đôn đốc nộp sau kiểm tra"},
        {"rate_col": "DC", "bq_val": 0.82, "type": "std", "task": "Kiểm tra đóng mã, giải thể"},
    ]

    for idx, team_id in enumerate(teams_list):
        row_num = 3 + idx
        team_name = [t["name"] for t in seed_data.DEMO_TEAMS if t["code"] == team_id][0]

        cell_name = ws_kpi_month.cell(row_num, 1, team_name)
        cell_name.font = kpi_bold_font
        cell_name.border = kpi_border
        cell_name.fill = subtle_fill
        ws_kpi_month.row_dimensions[row_num].height = 24

        # Specific formula for Indicator 2: Thực thu / Ước thu (Cols E, F, G, H)
        sh_thu = get_sh_name("Theo dõi số thu ngân sách")
        if sh_thu:
            last_r = wb[sh_thu].max_row
            ws_kpi_month[f"E{row_num}"] = f"=SUMIF('{sh_thu}'!C2:C{last_r}, \"{team_id}\", '{sh_thu}'!H2:H{last_r})"
            ws_kpi_month[f"F{row_num}"] = f"=SUMIF('{sh_thu}'!C2:C{last_r}, \"{team_id}\", '{sh_thu}'!G2:G{last_r})"
        else:
            ws_kpi_month[f"E{row_num}"] = 0
            ws_kpi_month[f"F{row_num}"] = 0
        ws_kpi_month[f"G{row_num}"] = f"=IF(F{row_num}>0, ABS(E{row_num}-F{row_num})/F{row_num}, 0)"
        ws_kpi_month[f"H{row_num}"] = f"=IF(G{row_num}<=0.03, 1, IF(G{row_num}<=0.05, 0.5, IF(G{row_num}<=0.1, 0, IF(G{row_num}<=0.15, -0.5, -1))))"

        for col_i in ["E", "F"]:
            c = ws_kpi_month[f"{col_i}{row_num}"]
            c.number_format = "#,##0"
            c.font = kpi_body_font
            c.border = kpi_border
            c.alignment = Alignment(horizontal="right")
        ws_kpi_month[f"G{row_num}"].number_format = "0.0%"
        ws_kpi_month[f"G{row_num}"].font = kpi_body_font
        ws_kpi_month[f"G{row_num}"].border = kpi_border
        ws_kpi_month[f"G{row_num}"].alignment = Alignment(horizontal="right")
        ws_kpi_month[f"H{row_num}"].number_format = "0.0"
        ws_kpi_month[f"H{row_num}"].font = kpi_bold_font
        ws_kpi_month[f"H{row_num}"].border = kpi_border
        ws_kpi_month[f"H{row_num}"].fill = kpi_score_fill
        ws_kpi_month[f"H{row_num}"].alignment = Alignment(horizontal="right")

        # Standard indicators
        for ind in month_indicators_map:
            r_col = ind["rate_col"]
            idx_num = column_index_from_string(r_col)
            bq_col = get_column_letter(idx_num + 1)
            sc_col = get_column_letter(idx_num + 2)

            sh_name = get_sh_name(ind["task"])
            if sh_name:
                task_ws = wb[sh_name]
                last_row = task_ws.max_row
                rate_formula = f"=IF(SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!G2:G{last_row})>0, SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!H2:H{last_row}) / SUMIF('{sh_name}'!C2:C{last_row}, \"{team_id}\", '{sh_name}'!G2:G{last_row}), 0)"
            else:
                rate_formula = 0

            c_rate = ws_kpi_month[f"{r_col}{row_num}"]
            c_rate.value = rate_formula
            c_rate.number_format = "0.0%"
            c_rate.font = kpi_body_font
            c_rate.border = kpi_border
            c_rate.alignment = Alignment(horizontal="right")

            c_bq = ws_kpi_month[f"{bq_col}{row_num}"]
            c_bq.value = ind["bq_val"]
            c_bq.number_format = "0.0%"
            c_bq.font = kpi_body_font
            c_bq.border = kpi_border
            c_bq.fill = kpi_avg_fill
            c_bq.alignment = Alignment(horizontal="right")

            c_score = ws_kpi_month[f"{sc_col}{row_num}"]
            if ind["type"] == "std_revenue":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, 1, IF({r_col}{row_num}>={bq_col}{row_num}, 0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, 0, IF({r_col}{row_num}>={bq_col}{row_num}-0.1, -0.5, -1))))"
            elif ind["type"] == "debt_growth":
                score_formula = f"=IF({r_col}{row_num}>=0.05, -1, IF({r_col}{row_num}>=0, -0.5, IF({r_col}{row_num}>=-0.05, 0.5, 1)))"
            elif ind["type"] == "debt_ratio":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.02, -1, IF({r_col}{row_num}>={bq_col}{row_num}, -0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.02, 0.5, 1)))"
            elif ind["type"] == "overdue":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, -1, IF({r_col}{row_num}>={bq_col}{row_num}, -0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, 0.5, 1)))"
            elif ind["type"] == "sat":
                score_formula = f"=IF({r_col}{row_num}>=0.95, 1, IF({r_col}{row_num}>=0.9, 0.5, IF({r_col}{row_num}>=0.85, -0.5, -1)))"
            elif ind["type"] == "std":
                score_formula = f"=IF({r_col}{row_num}>={bq_col}{row_num}+0.05, 1, IF({r_col}{row_num}>={bq_col}{row_num}, 0.5, IF({r_col}{row_num}>={bq_col}{row_num}-0.05, -0.5, -1)))"
            else:
                score_formula = 0

            c_score.value = score_formula
            c_score.number_format = "0.0"
            c_score.font = kpi_bold_font
            c_score.border = kpi_border
            c_score.fill = kpi_score_fill
            c_score.alignment = Alignment(horizontal="right")

        # Category Averages
        # L: =AVERAGE(D, H, K)
        ws_kpi_month[f"L{row_num}"] = f"=AVERAGE(D{row_num}, H{row_num}, K{row_num})"
        # AN: =AVERAGE(O, R, U, X, AA, AD, AG, AJ, AM)
        ws_kpi_month[f"AN{row_num}"] = f"=AVERAGE(O{row_num}, R{row_num}, U{row_num}, X{row_num}, AA{row_num}, AD{row_num}, AG{row_num}, AJ{row_num}, AM{row_num})"
        # AX: =AVERAGE(AQ, AT, AW)
        ws_kpi_month[f"AX{row_num}"] = f"=AVERAGE(AQ{row_num}, AT{row_num}, AW{row_num})"
        # BZ: =AVERAGE(BA, BD, BG, BJ, BM, BP, BS, BV, BY)
        ws_kpi_month[f"BZ{row_num}"] = f"=AVERAGE(BA{row_num}, BD{row_num}, BG{row_num}, BJ{row_num}, BM{row_num}, BP{row_num}, BS{row_num}, BV{row_num}, BY{row_num})"
        # CS: =AVERAGE(CC, CF, CI, CL, CO, CR)
        ws_kpi_month[f"CS{row_num}"] = f"=AVERAGE(CC{row_num}, CF{row_num}, CI{row_num}, CL{row_num}, CO{row_num}, CR{row_num})"
        # DF: =AVERAGE(CV, CY, DB, DE)
        ws_kpi_month[f"DF{row_num}"] = f"=AVERAGE(CV{row_num}, CY{row_num}, DB{row_num}, DE{row_num})"
        # DG (Total Score): =SUM(L, AN, AX, BZ, CS, DF)
        ws_kpi_month[f"DG{row_num}"] = f"=SUM(L{row_num}, AN{row_num}, AX{row_num}, BZ{row_num}, CS{row_num}, DF{row_num})"
        # DH (Rank): =RANK(DG, DG$3:DG$10)
        ws_kpi_month[f"DH{row_num}"] = f"=RANK(DG{row_num}, DG$3:DG$10)"

        for g_col in ["L", "AN", "AX", "BZ", "CS", "DF", "DG", "DH"]:
            cell = ws_kpi_month[f"{g_col}{row_num}"]
            cell.fill = kpi_green_fill
            cell.font = kpi_bold_font
            cell.border = kpi_border
            cell.number_format = "0.00" if g_col != "DH" else "0"
            cell.alignment = Alignment(horizontal="right")

    # Column widths for monthly KPI
    ws_kpi_month.column_dimensions["A"].width = 30
    for c_idx in range(2, 114):
        col_letter = get_column_letter(c_idx)
        ws_kpi_month.column_dimensions[col_letter].width = 11

    # Re-apply validations on "Du lieu" with the actual row count
    ws_dest.data_validations.dataValidation.clear()
    max_row_idx = row_idx - 1
    
    tasks_len = len(seed_data.DEMO_TASKS)
    officers = []
    for team_id, staff_list in seed_data.DEMO_STAFF.items():
        for s in staff_list:
            if s["name"] not in officers:
                officers.append(s["name"])
    officers_len = len(officers)
    teams_len = len(seed_data.DEMO_TEAMS)
    
    validations_new = [
        (DataValidation(type="list", formula1=f"='Danh muc'!$A$2:$A${tasks_len + 1}", allow_blank=True), f"B2:B{max_row_idx}", "Chọn tên nhiệm vụ trong danh mục."),
        (DataValidation(type="list", formula1=f"='Danh muc'!$C$2:$C${officers_len + 1}", allow_blank=True), f"C2:C{max_row_idx}", "Chọn tên cán bộ trong danh mục."),
        (DataValidation(type="list", formula1=f"='Danh muc'!$D$2:$D$21", allow_blank=True), f"D2:D{max_row_idx}", "Chọn địa bàn xã cũ trong danh mục."),
        (DataValidation(type="list", formula1=f"='Danh muc'!$E$2:$E$5", allow_blank=True), f"E2:E{max_row_idx}", "Chọn địa bàn xã mới trong danh mục."),
        (DataValidation(type="list", formula1=f"='Danh muc'!$F$2:$F${teams_len + 1}", allow_blank=True), f"F2:F{max_row_idx}", "Chọn tổ quản lý trong danh mục."),
        (DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True), f"I2:J{max_row_idx}", "Chỉ nhập số nguyên không âm."),
        (DataValidation(type="date", operator="greaterThanOrEqual", formula1="DATE(2024,1,1)", allow_blank=True), f"K2:K{max_row_idx}", "Nhập ngày hợp lệ, nên dùng yyyy-mm-dd."),
    ]
    for val_obj, target, message in validations_new:
        val_obj.errorTitle = "Dữ liệu không hợp lệ"
        val_obj.error = message
        ws_dest.add_data_validation(val_obj)
        val_obj.add(target)

    # Save output file
    try:
        wb.save(OUTPUT_FILE)
        print(f"ĐÃ TẠO THÀNH CÔNG FILE TỔNG HỢP: {OUTPUT_FILE}")
    except PermissionError:
        print("\n" + "="*80)
        print("LỖI: KHÔNG THỂ GHI ĐÈ FILE EXCEL!")
        print(f"Vui lòng ĐÓNG file '{OUTPUT_FILE.name}' trong Microsoft Excel trước khi chạy script!")
        print("="*80 + "\n")
        sys.exit(1)


if __name__ == "__main__":
    main()
