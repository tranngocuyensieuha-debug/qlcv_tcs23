import sys
from pathlib import Path
from openpyxl import load_workbook

def main():
    f = Path(r"D:\TCS23_QLCV-20260826T011945Z-1-001\TCS23_QLCV\TCS23_Quản lý công việc\Final W5M7_Bộ tiêu chí đánh giá các chỉ tiêu.xlsx")
    out_path = Path(__file__).resolve().parent / "inspect_output.txt"
    
    with open(out_path, "w", encoding="utf-8") as out:
        out.write(f"Đường dẫn file: {f}\n")
        out.write(f"File tồn tại: {f.exists()}\n")
        if not f.exists():
            return
            
        wb = load_workbook(f, data_only=True)
        out.write(f"Các sheet có trong file: {wb.sheetnames}\n")
        for s_name in wb.sheetnames:
            ws = wb[s_name]
            out.write(f"\n=== Sheet '{s_name}' ({ws.max_row} dòng, {ws.max_column} cột) ===\n")
            printed_count = 0
            for r in range(1, ws.max_row + 1):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                if any(val is not None for val in row_vals):
                    row_str = [str(v) if v is not None else "" for v in row_vals[:12]]
                    out.write(f"Dòng {r}: {row_str}\n")
                    printed_count += 1
                    if printed_count >= 150:
                        out.write("... (đã ẩn bớt các dòng sau) ...\n")
                        break
                        
    print(f"ĐÃ LƯU KẾT QUẢ VÀO FILE: {out_path.name}")

if __name__ == "__main__":
    main()
