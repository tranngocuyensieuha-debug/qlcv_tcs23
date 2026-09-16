import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

# Explicit base directories
TARGET_DIR = Path(r"d:\TCS23_QLCV-20260826T011945Z-1-001\TCS23_QLCV")
SCRIPTS_DIR = TARGET_DIR / "quan ly cong viec" / "quan ly cong viec" / "app" / "scripts"
OUT_FILE = SCRIPTS_DIR / "inspect_monthly_output.txt"

lines = []

def log(msg):
    lines.append(str(msg))
    print(msg)

log("=" * 80)
log(f"TARGET DIRECTORY: {TARGET_DIR}")
log(f"OUTPUT FILE: {OUT_FILE}")
log("=" * 80)

# 1. READ "TCS23 - CV Chỉ tiêu.docx"
docx_path = TARGET_DIR / "TCS23 - CV Chỉ tiêu.docx"
log(f"\n1. INSPECTING DOCX: {docx_path}")
if docx_path.exists():
    try:
        with zipfile.ZipFile(docx_path) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            
            log("\n--- NỘI DUNG VĂN BẢN (PARAGRAPHS & TABLES) ---")
            for elem in tree.iter():
                # Paragraph
                if elem.tag.endswith("}p"):
                    texts = [n.text for n in elem.iter() if n.tag.endswith("}t") and n.text]
                    if texts:
                        text_str = "".join(texts).strip()
                        if text_str:
                            log("P: " + text_str)
                # Table row
                elif elem.tag.endswith("}tr"):
                    row_texts = []
                    for tc in elem.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"):
                        cell_texts = [n.text for n in tc.iter() if n.tag.endswith("}t") and n.text]
                        row_texts.append("".join(cell_texts).strip())
                    if any(row_texts):
                        log("ROW: " + " | ".join(row_texts))
    except Exception as e:
        log(f"Error reading docx: {e}")
else:
    log(f"File not found: {docx_path}")

# 2. READ "Bộ tiêu chí.xlsx"
xlsx_path = TARGET_DIR / "Bộ tiêu chí.xlsx"
log(f"\n2. INSPECTING XLSX: {xlsx_path}")
if xlsx_path.exists():
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        log(f"Sheet names: {wb.sheetnames}")
        for name in wb.sheetnames:
            ws = wb[name]
            log(f"\n--- SHEET: {name} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
            for r in range(1, min(ws.max_row + 1, 150)):
                row_vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else "" for c in range(1, min(ws.max_column + 1, 26))]
                if any(v for v in row_vals):
                    log(f"R{r:02d}: " + " | ".join(row_vals))
    except Exception as e:
        log(f"Error reading with openpyxl: {e}")
else:
    log(f"File not found: {xlsx_path}")

try:
    with open(OUT_FILE, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    log(f"\n>>> ĐÃ GHI THÀNH CÔNG DỮ LIỆU RA: {OUT_FILE}")
except Exception as e:
    log(f"Error writing output file: {e}")
